"""Read an uploaded file from stdin. Never executes macros or follows external links."""
import base64, csv, hashlib, io, json, os, sys, zipfile
sys.stdout.reconfigure(encoding='utf-8')
data = json.load(sys.stdin)
blob = base64.b64decode(data['base64'], validate=True)
if len(blob) > 5 * 1024 * 1024:
    raise ValueError('File exceeds 5 MB')
ext = os.path.splitext(data['name'])[1].lower()
if ext in ('.docx', '.xlsx'):
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        if sum(x.file_size for x in z.infolist()) > 30 * 1024 * 1024:
            raise ValueError('Archive expands beyond 30 MB')
        if any('vbaProject' in x.filename for x in z.infolist()):
            raise ValueError('Macros are not accepted')
result = {'fileChecksum': hashlib.sha256(blob).hexdigest(), 'rows': None, 'pages': []}
if ext == '.csv':
    result['rows'] = list(csv.DictReader(io.StringIO(blob.decode('utf-8-sig'))))
elif ext == '.xlsx':
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True, keep_links=False)
    rows = iter(book.active.iter_rows(values_only=True))
    headers = [str(x or '') for x in next(rows)]
    result['rows'] = [dict(zip(headers, [str(x) if x is not None else '' for x in row])) for i, row in enumerate(rows) if i < 500 and any(x is not None for x in row)]
elif ext == '.docx':
    from docx import Document
    doc = Document(io.BytesIO(blob))
    result['pages'] = [{'locator': 'paragraph:' + str(i + 1), 'text': p.text} for i, p in enumerate(doc.paragraphs) if p.text.strip()]
    result['pages'] += [{'locator': 'table:' + str(i + 1), 'text': '\n'.join('\t'.join(c.text for c in r.cells) for r in t.rows)} for i, t in enumerate(doc.tables)]
elif ext == '.pdf':
    import pdfplumber
    with pdfplumber.open(io.BytesIO(blob)) as pdf:
        if len(pdf.pages) > 100:
            raise ValueError('More than 100 pages')
        result['pages'] = [{'locator': 'page:' + str(i + 1), 'text': p.extract_text() or ''} for i, p in enumerate(pdf.pages)]
else:
    raise ValueError('Unsupported format')
result['text'] = '\n'.join(p['locator'] + '\n' + p['text'] for p in result['pages'])
print(json.dumps(result, ensure_ascii=False))
