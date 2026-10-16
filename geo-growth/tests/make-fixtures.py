from pathlib import Path
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen.canvas import Canvas
folder=Path(__file__).resolve().parents[1]/'data'
folder.mkdir(exist_ok=True)
doc=Document()
doc.add_heading('Fixture Product Manual',0)
doc.add_paragraph('SKU DOC-1. Power: 6 kW. Synthetic test fixture only.')
table=doc.add_table(rows=1,cols=2)
table.cell(0,0).text='Material'
table.cell(0,1).text='Steel'
doc.save(folder/'manual-fixture.docx')
book=Workbook()
book.active.append(['sku','name','category','power'])
book.active.append(['XLS-1','Fixture Pump','pump','4 kW'])
book.save(folder/'catalog-fixture.xlsx')
canvas=Canvas(str(folder/'manual-fixture.pdf'))
canvas.drawString(50,780,'SKU PDF-1. Power: 7 kW. Synthetic extraction fixture.')
canvas.save()
(folder/'catalog-template.csv').write_text('sku,name,category,material,power\nNEW-1,Your product,centrifugal pump,Stainless steel 304,5 kW\n',encoding='utf-8')
